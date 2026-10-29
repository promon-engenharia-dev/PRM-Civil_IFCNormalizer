import openpyxl

from conftest import build_minimal_model
from ifc_normalizer import migrate, pipeline, rules, validation_pack

TIEBAR = {"quando": {"origem": "Tekla", "classe": ["IfcBeam", "IfcColumn"], "nome": {"em": ["CORRENTE", "TIRANTE"]}},
          "entao": {"classe": "IfcMember", "predefined_type": "TIEBAR"}}


def test_plain_language_texts():
    applies = validation_pack.describe_applies_to(TIEBAR["quando"])
    assert applies == ("Elementos do Tekla, classificados como viga (IfcBeam) ou pilar (IfcColumn), "
                       "quando o nome é CORRENTE ou TIRANTE.")
    assert validation_pack.describe_change(TIEBAR["entao"]) == \
        "Passa a ser barra estrutural (IfcMember), tipo barra de tração (TIEBAR)."

    copy_rule = {"propriedades": {"PRM-Parametros": {
        "PRM_Altura": {"copiar": "PRM_Dimensoes.Altura", "unidade_origem": "m"}}}}
    assert "Copia “Altura” (grupo PRM_Dimensoes) para PRM-Parametros › PRM_Altura (convertendo de metros" in \
        validation_pack.describe_change(copy_rule)

    cond = {"origem": "Revit", "propriedades": {"PRM_Construtibilidade.CodigoAtividade": {"vazio": False}}}
    assert "o parâmetro “CodigoAtividade” (grupo PRM_Construtibilidade) está preenchido" in \
        validation_pack.describe_applies_to(cond)


def test_technical_texts_kept_for_editors():
    assert validation_pack.describe_condition(TIEBAR["quando"]) == \
        "origem: Tekla; classe: IfcBeam, IfcColumn; nome em [CORRENTE, TIRANTE]"
    assert validation_pack.describe_action(TIEBAR["entao"]) == "classe = IfcMember; PredefinedType = TIEBAR"


def test_build_pack_from_simulation(tmp_path):
    f = build_minimal_model()
    f.by_type("IfcBuildingElementProxy")[0].Name = "NUT"
    source = tmp_path / "ZTO-00-ST3-TESTE.ifc"
    f.write(str(source))
    rs = rules.load_normalization_rules()
    out = tmp_path / "sim"
    out.mkdir()
    summary = pipeline.normalize_one(source, out, rs, migrate.load_rules(), include_proposed=True,
                                     run_validation=False)
    assert summary["status"] == "ok", summary

    pack = validation_pack.build(rs, sorted(out.glob("*.normalizacao.xlsx")), tmp_path / "pacote.xlsx")
    wb = openpyxl.load_workbook(pack)
    assert wb.sheetnames[:2] == ["LEIA-ME", "Regras_Validacao"] and "LISTAS" in wb.sheetnames

    ws = wb["Regras_Validacao"]
    header = [c.value for c in ws[5]][1:]
    for col in ("Aplica-se a", "O que muda", "Exemplo real (antes → depois)", "Perguntas para validar",
                *validation_pack.DECISION_COLUMNS):
        assert col in header
    ids = [ws.cell(row=r, column=3).value for r in range(6, ws.max_row + 1)]
    assert [i for i in ids if i] == [r["id"] for r in sorted(
        rs.rules, key=lambda r: [x.get("tema") for x in rs.rules].index(r.get("tema")))]
    assert ws.sheet_view.showGridLines is False
    assert ws.data_validations.dataValidation

    readme = " ".join(str(c.value) for row in wb["LEIA-ME"].iter_rows() for c in row if c.value)
    assert "Glossário" in readme and "PredefinedType" in readme

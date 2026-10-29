"""Pacote de validação das regras com o BEP (planilha no padrão Promon).

Feito para quem não acompanhou o desenvolvimento: aba LEIA-ME com instruções e glossário,
regras explicadas em linguagem direta com exemplo real de antes → depois e perguntas para
validar, e a revisão das interpretações geométricas. O detalhe técnico (formato do JSON)
fica nas últimas colunas, só como referência.

Entrada: JSON de regras + relatórios *.normalizacao.xlsx de uma simulação
(`normalize --incluir-propostas`).
"""

from __future__ import annotations

import datetime as dt
import math
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .profile import normalize_name
from .rules import RuleSet, uses_geometry

# Padrão de planilha Promon (manual de formatação, item 7.2)
AZUL, AZUL_ESC, CINZA, CINZA_02 = "004976", "01426A", "333F48", "5B6770"
LARANJA_ESC = "BE531C"
ZEBRA, REFERENCIA, INFO = "EBF0F5", "DDE6EF", "F4F7F9"
FONT = "Tahoma"
THIN = Side(style="thin", color="E2E8EC")

DECISOES = ["Aprovar", "Ajustar", "Rejeitar", "Aguardar"]
DECISION_COLUMNS = ["Decisão", "Comentário ou ajuste", "Referência no BEP", "Responsável", "Data"]

# ------------------------------------------------- vocabulário em português
# Só apresentação: traduz nomes IFC para a leitura de quem valida.
CLASSES = {
    "IfcBeam": "viga", "IfcColumn": "pilar", "IfcMember": "barra estrutural", "IfcPlate": "chapa",
    "IfcSlab": "laje", "IfcWall": "parede", "IfcFooting": "fundação", "IfcCovering": "revestimento",
    "IfcMechanicalFastener": "fixador mecânico", "IfcBuildingElementProxy": "elemento genérico (proxy)",
    "IfcElementAssembly": "conjunto montado", "IfcSanitaryTerminal": "aparelho sanitário",
    "IfcCourse": "camada (pavimento/lastro)", "IfcPavement": "pavimento", "IfcKerb": "guia (meio-fio)",
    "IfcRail": "trilho", "IfcGeographicElement": "elemento geográfico", "IfcPipeSegment": "tubo",
    "IfcEarthworksFill": "aterro", "IfcEarthworksCut": "corte",
}
PREDEFINED = {
    "TIEBAR": "barra de tração", "BRACE": "contraventamento", "PURLIN": "terça", "MEMBER": "barra",
    "BEAM": "viga", "COLUMN": "pilar", "ANCHORBOLT": "chumbador", "BOLT": "parafuso",
    "PAVEMENT": "revestimento", "BALLASTBED": "lastro", "RAIL": "trilho", "TERRAIN": "terreno",
    "SANITARYFOUNTAIN": "fonte sanitária (bebedouro)", "NOTDEFINED": "não definido",
}
DISCIPLINAS = {
    "ST3": "estrutura metálica (ST3)", "ST4": "estrutura de concreto (ST4)", "S": "estrutura (S)",
    "AQ": "arquitetura (AQ)", "A": "arquitetura (A)", "C": "civil (C)", "IH": "infraestrutura (IH)",
}
CAMPOS = {
    "nome": "o nome", "descricao": "a descrição", "object_type": "o ObjectType", "tipo_nome": "o nome do tipo",
    "layer": "a camada (layer)", "arquivo": "o nome do arquivo", "predefined_type": "o PredefinedType",
}
ORIENTACOES = {
    "vertical": "lineares e verticais", "horizontal": "lineares e horizontais", "inclinado": "lineares e inclinadas",
    "plano_horizontal": "planas e horizontais", "plano_vertical": "planas e verticais",
    "plano_inclinado": "planas e inclinadas", "compacto": "compactas", "sem_geometria": "sem geometria",
}
CAMPOS_ALTERADOS = {"classe": "Classe IFC", "PredefinedType": "Tipo predefinido", "ObjectType": "ObjectType",
                    "Name": "Nome", "Description": "Descrição", "ElementType": "Tipo do elemento"}

GLOSSARIO = [
    ("IFC", "Formato aberto da buildingSMART para troca de modelos BIM. A versão-alvo é IFC 4.3 (IFC4X3_ADD2)."),
    ("Classe IFC", "O que o elemento é: viga (IfcBeam), pilar (IfcColumn), laje (IfcSlab), trilho (IfcRail)… "
                   "Define como o elemento é reconhecido por qualquer programa."),
    ("Tipo predefinido (PredefinedType)", "Detalhe da classe, escolhido de uma lista fixa do IFC. "
                   "Ex.: barra estrutural (IfcMember) do tipo contraventamento (BRACE) ou terça (PURLIN)."),
    ("USERDEFINED / ObjectType", "Quando a lista do IFC não tem o tipo certo: o tipo predefinido fica "
                   "USERDEFINED e o nome do tipo vai em ObjectType (ex.: PORCA)."),
    ("Elemento genérico (proxy)", "IfcBuildingElementProxy: elemento sem classe definida. É o que mais "
                   "precisa de regra, porque não informa o que a peça é."),
    ("Pset (conjunto de propriedades)", "Grupo de parâmetros do elemento. O padrão Promon é o Pset "
                   "PRM-Parametros (código de atividade, EAP, dimensões…)."),
    ("Tipo (IfcTypeObject)", "Definição compartilhada por vários elementos iguais (ex.: um perfil do Tekla). "
                   "Quando a classe do elemento muda, o tipo precisa acompanhar."),
    ("GlobalId", "Identificador único e permanente do elemento. A ferramenta nunca o altera; serve para "
                   "localizar a peça no visualizador (Navisworks, BIMcollab, Solibri…)."),
    ("Regra", "Instrução 'se o elemento atende a X, então mude Y'. Fica no arquivo rules/normalizacao.json, "
              "ao lado do executável."),
    ("Proposta / aprovada", "Regra proposta ainda não foi validada: só roda em simulação. "
                            "Regra aprovada roda sempre."),
    ("Pendência", "Elemento que nenhuma regra resolveu. Aparece no relatório; a ferramenta nunca escolhe "
                  "a classe por suposição."),
    ("Heurística geométrica", "Classificação pela forma da peça (vertical, horizontal, plana). É interpretação "
                              "e por isso fica em teste até ser validada."),
]


def class_label(c: str) -> str:
    return f"{CLASSES[c]} ({c})" if c in CLASSES else c


def _join(items: list[str], last: str = "ou") -> str:
    items = [str(i) for i in items]
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + f" {last} " + items[-1]


# --------------------------------------------------------- texto das regras

def _phrase_matcher(field: str, value) -> str:
    label = CAMPOS.get(field, field)
    if isinstance(value, dict):
        op, arg = next(iter(value.items()))
        if op == "em":
            return f"{label} é {_join(arg)}"
        if op == "contem":
            return f"{label} contém “{arg}”"
        if op == "comeca_com":
            return f"{label} começa com “{arg}”"
        if op == "regex":
            return f"{label} segue o padrão {arg}"
        if op == "vazio":
            return f"{label} está {'vazio' if arg else 'preenchido'}"
        return f"{label} é {arg}"
    return f"{label} é {value}"


def _phrase_property(prop: str, value) -> str:
    pset, name = prop.split(".", 1)
    if isinstance(value, dict) and "vazio" in value:
        return f"o parâmetro “{name}” (grupo {pset}) está {'vazio' if value['vazio'] else 'preenchido'}"
    if isinstance(value, dict) and "em" in value:
        return f"o parâmetro “{name}” (grupo {pset}) é {_join(value['em'])}"
    return f"o parâmetro “{name}” (grupo {pset}) é {value if not isinstance(value, dict) else next(iter(value.values()))}"


def _phrase_conditions(cond: dict) -> list[str]:
    parts = []
    for key, value in cond.items():
        if key in ("origem", "disciplina", "classe"):
            continue
        if key == "ou":
            parts.append("(" + " ou ".join("; ".join(_phrase_conditions(c)) for c in value) + ")")
        elif key == "propriedades":
            parts += [_phrase_property(p, v) for p, v in value.items()]
        elif key == "geometria":
            opts = value if isinstance(value, str) else value.get("em", [])
            opts = [opts] if isinstance(opts, str) else opts
            parts.append("a forma da peça é " + _join([ORIENTACOES.get(o, o) for o in opts]))
        else:
            parts.append(_phrase_matcher(key, value))
    return parts


def describe_applies_to(cond: dict) -> str:
    """'Elementos do Tekla, estrutura (S), classificados como viga ou pilar, quando o nome é NUT'."""
    def as_list(v):
        return v if isinstance(v, list) else [v]

    subject = "Elementos"
    if "origem" in cond:
        subject += " do " + _join(as_list(cond["origem"]))
    if "disciplina" in cond:
        subject += " – " + _join([DISCIPLINAS.get(d, d) for d in as_list(cond["disciplina"])])
    if "classe" in cond:
        subject += ", classificados como " + _join([class_label(c) for c in as_list(cond["classe"])])
    conditions = _phrase_conditions(cond)
    return subject + (", quando " + " e ".join(conditions) if conditions else "") + "."


def describe_change(action: dict) -> str:
    lines = []
    if "classe" in action:
        line = f"Passa a ser {class_label(action['classe'])}"
        pdt = action.get("predefined_type")
        if pdt == "USERDEFINED":
            line += f", tipo definido pelo usuário “{action.get('object_type', '')}”"
        elif pdt:
            line += f", tipo {PREDEFINED.get(pdt, pdt)} ({pdt})"
        lines.append(line + ".")
    elif "predefined_type" in action:
        pdt = action["predefined_type"]
        lines.append(f"Tipo predefinido passa a ser {PREDEFINED.get(pdt, pdt)} ({pdt}).")
    if "object_type" in action and action.get("predefined_type") != "USERDEFINED":
        lines.append(f"ObjectType passa a ser “{action['object_type']}”.")
    if "nome" in action:
        lines.append(f"Nome passa a ser “{action['nome']}”.")
    if "descricao" in action:
        lines.append(f"Descrição passa a ser “{action['descricao']}”.")
    for pset, props in (action.get("propriedades") or {}).items():
        for prop, spec in props.items():
            if isinstance(spec, dict):
                src = spec["copiar"].split(".", 1)
                unit = " (convertendo de metros para a unidade do arquivo)" if spec.get("unidade_origem") == "m" else ""
                lines.append(f"Copia “{src[1]}” (grupo {src[0]}) para {pset} › {prop}{unit}.")
            else:
                lines.append(f"Preenche {pset} › {prop} = {spec}.")
    return "\n".join(lines)


# Formato técnico (para quem edita o JSON)

def _matcher(value) -> str:
    if isinstance(value, dict):
        op, arg = next(iter(value.items()))
        if op == "em":
            return "em [" + ", ".join(map(str, arg)) + "]"
        if op == "vazio":
            return "vazio" if arg else "preenchido"
        return f"{op.replace('_', ' ')} '{arg}'"
    return f"= '{value}'"


def describe_condition(cond: dict) -> str:
    parts = []
    for key, value in cond.items():
        if key == "ou":
            parts.append("(" + " OU ".join(describe_condition(c) for c in value) + ")")
        elif key == "propriedades":
            parts += [f"{prop} {_matcher(m)}" for prop, m in value.items()]
        elif key in ("origem", "disciplina", "classe") or (key == "geometria" and isinstance(value, str)):
            parts.append(f"{key}: {', '.join(value) if isinstance(value, list) else value}")
        else:
            parts.append(f"{key} {_matcher(value)}")
    return "; ".join(parts)


def describe_action(action: dict) -> str:
    labels = {"classe": "classe", "predefined_type": "PredefinedType", "object_type": "ObjectType",
              "nome": "Name", "descricao": "Description"}
    parts = [f"{label} = {action[key]}" for key, label in labels.items() if key in action]
    for pset, props in (action.get("propriedades") or {}).items():
        for prop, spec in props.items():
            if isinstance(spec, dict):
                unit = f" ({spec['unidade_origem']})" if spec.get("unidade_origem") else ""
                parts.append(f"{pset}.{prop} ← {spec['copiar']}{unit}")
            else:
                parts.append(f"{pset}.{prop} = {spec}")
    return "; ".join(parts)


# ----------------------------------------------------------------- dados

def _value(v) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)) or str(v).strip() == "":
        return "(vazio)"
    if isinstance(v, float):  # 0.5750000000000001 -> 0,575
        return f"{v:.3f}".rstrip("0").rstrip(".").replace(".", ",")
    text = str(v)
    return text if len(text) <= 60 else text[:57] + "…"


def collect(reports: list[Path]) -> tuple[dict, pd.DataFrame]:
    """Uso por regra (somado), exemplo real de antes → depois e revisão geométrica."""
    usage: dict[str, dict] = defaultdict(lambda: {"casados": 0, "alteracoes": 0, "arquivos": 0,
                                                    "nomes": Counter(), "por_nome": {}, "exemplo": None})
    geo_frames = []
    for path in reports:
        xls = pd.ExcelFile(path)
        name = path.name.replace(".normalizacao.xlsx", ".ifc")
        for _, row in pd.read_excel(xls, "Uso_Regras").iterrows():
            u = usage[row["regra"]]
            u["casados"] += int(row["elementos_casados"])
            u["alteracoes"] += int(row["alteracoes"])
            u["arquivos"] += int(row["elementos_casados"] > 0)
        changes = pd.read_excel(xls, "Alteracoes")
        for (rid, guid), rows in changes.groupby(["regra", "global_id"], sort=False):
            u = usage[rid]
            if rows["classe"].iloc[0].endswith("Type"):
                continue  # o exemplo é de elemento; alteração de tipo aparece no relatório de cada arquivo
            nome = rows["nome"].iloc[0]
            key = normalize_name(nome) if isinstance(nome, str) and nome else "(sem nome)"
            u["nomes"][key] += 1
            u["por_nome"].setdefault(key, {"arquivo": name, "nome": nome if isinstance(nome, str) else "(sem nome)",
                                           "linhas": [
                f"{CAMPOS_ALTERADOS.get(r.campo, r.campo.replace('.', ' › '))}: {_value(r.antes)} → {_value(r.depois)}"
                for r in rows.itertuples()]})
        if "Revisao_Geometrica" in xls.sheet_names:
            g = pd.read_excel(xls, "Revisao_Geometrica")
            g["arquivo"] = name
            geo_frames.append(g)
    # Exemplo: peça com o nome mais frequente da regra, evitando peças sem nome quando houver outra.
    for u in usage.values():
        named = [(n, q) for n, q in u["nomes"].most_common() if n.lower() not in ("(sem nome)", "unnamed")]
        best = (named or u["nomes"].most_common())[:1]
        u["exemplo"] = u["por_nome"][best[0][0]] if best else None
    return usage, (pd.concat(geo_frames) if geo_frames else pd.DataFrame())


def geometry_groups(geo: pd.DataFrame) -> pd.DataFrame:
    """Interpretações que pedem decisão: aplicadas pela geometria e inclinadas sem classificação."""
    if geo.empty:
        return geo
    geo = geo.copy()
    geo["nome_normalizado"] = geo["nome"].fillna("").map(normalize_name)
    unresolved = geo["predefined_final"].fillna("").isin(["", "NOTDEFINED"])
    inclined = geo["orientacao"].isin(["inclinado", "plano_inclinado"])
    relevant = geo[(geo["situacao"] == "aplicada") | (inclined & unresolved)]
    out = (
        relevant.groupby(["nome_normalizado", "classe_original", "orientacao", "situacao", "classe_final",
                          "predefined_final"], dropna=False)
        .agg(qtd=("global_id", "size"), angulo_medio=("angulo_vertical_graus", "mean"),
             comprimento_mediano_m=("comprimento_m", "median"), arquivos=("arquivo", "nunique"),
             exemplo_arquivo=("arquivo", "first"), exemplo_global_id=("global_id", "first"))
        .reset_index()
        .sort_values("qtd", ascending=False)
    )
    out["angulo_medio"] = out["angulo_medio"].round(0)
    out["comprimento_mediano_m"] = out["comprimento_mediano_m"].round(2)
    return out


def geometry_texts(row) -> tuple[str, str]:
    """(o que a ferramenta faria, pergunta) para um grupo da revisão geométrica."""
    nome = row.nome_normalizado or "(sem nome)"
    forma = ORIENTACOES.get(row.orientacao, row.orientacao)
    medida = f"comprimento típico {row.comprimento_mediano_m:.2f} m".replace(".", ",") \
        if pd.notna(row.comprimento_mediano_m) else "sem medida"
    pdt = row.predefined_final if isinstance(row.predefined_final, str) else ""
    if row.situacao == "aplicada":
        destino = class_label(row.classe_final) + (f", tipo {PREDEFINED.get(pdt, pdt)}" if pdt else "")
        acao = f"Passaria a {destino}."
        pergunta = (f"As {row.qtd} peças “{nome}” {forma} ({medida}) devem mesmo virar "
                    f"{CLASSES.get(row.classe_final, row.classe_final)}?")
    else:
        acao = "Nenhuma regra: fica como está e entra como pendência."
        angulo = f"~{int(row.angulo_medio)}° da vertical" if pd.notna(row.angulo_medio) else ""
        pergunta = f"Como classificar as {row.qtd} peças “{nome}” {forma} ({angulo}, {medida})?"
    return acao, pergunta


# ----------------------------------------------------------------- estilo

def _fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def _font(size=9, bold=False, color=CINZA, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def _band(ws, row: int, text: str, ncols: int, fill: str, size: int, color="FFFFFF", height=None, center=False):
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=ncols + 1)
    c = ws.cell(row=row, column=2, value=text)
    c.font = _font(size, True, color)
    c.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True, indent=1)
    for col in range(2, ncols + 2):
        ws.cell(row=row, column=col).fill = _fill(fill)
    if height:
        ws.row_dimensions[row].height = height


def _row_height(values: list, widths: list[int]) -> float:
    """Altura estimada para texto quebrado (Tahoma 9 ≈ 1 caractere por unidade de largura)."""
    lines = 1
    for v, w in zip(values, widths):
        if v is None:
            continue
        text_lines = str(v).split("\n")
        lines = max(lines, sum(max(1, math.ceil(len(t) / max(w - 1, 1))) for t in text_lines))
    return min(12.5 * lines + 4, 300)


def _header(ws, row: int, headers: list[str], widths: list[int], decision_from: int | None):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=j + 2, value=h)
        in_decision = decision_from is not None and decision_from <= j < decision_from + len(DECISION_COLUMNS)
        c.font = _font(9, True, "FFFFFF")
        c.fill = _fill(LARANJA_ESC if in_decision else AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j + 2)].width = widths[j]
    ws.row_dimensions[row].height = 32


def _data_row(ws, row: int, values: list, widths: list[int], zebra: bool, numeric: set[int],
              decision_from: int | None, technical_from: int | None):
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=j + 2, value=v)
        technical = technical_from is not None and j >= technical_from
        c.font = _font(8 if technical else 9, False, CINZA_02 if technical else (AZUL if j in numeric else CINZA))
        c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="right" if j in numeric else "left")
        c.border = Border(bottom=THIN)
        if zebra:
            c.fill = _fill(ZEBRA)
        if decision_from is not None and decision_from <= j < decision_from + len(DECISION_COLUMNS):
            c.border = Border(bottom=THIN, left=THIN, right=THIN)
    ws.row_dimensions[row].height = _row_height(values, widths)


def _sheet_top(ws, title: str, subtitle: str, ncols: int, instructions: str):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    _band(ws, 1, title, ncols, AZUL, 14, height=28, center=True)
    _band(ws, 2, subtitle, ncols, AZUL_ESC, 10, height=20)
    _band(ws, 3, instructions, ncols, REFERENCIA, 9, color=AZUL, height=48)


def _decision_list(ws, n: int) -> DataValidation:
    dv = DataValidation(type="list", formula1=f"=LISTAS!$A$2:$A${n + 1}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Decisão", error="Escolha um valor da lista.")
    ws.add_data_validation(dv)
    return dv


# ---------------------------------------------------------------- abas

def _readme(wb: Workbook, ruleset: RuleSet, n_reports: int, today: str):
    ws = wb.create_sheet("LEIA-ME", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 100
    _band(ws, 1, "NORMALIZADOR IFC – VALIDAÇÃO DAS REGRAS", 2, AZUL, 14, height=28, center=True)
    _band(ws, 2, f"Regras versão {ruleset.version} · simulação em {n_reports} arquivo(s) · gerado em {today}",
          2, AZUL_ESC, 10, height=20)

    row = 4

    def section(title: str):
        nonlocal row
        row += 1
        _band(ws, row, title, 2, REFERENCIA, 11, color=AZUL, height=22)
        row += 1

    def text(t: str, bold_label: str | None = None):
        nonlocal row
        if bold_label:
            a = ws.cell(row=row, column=2, value=bold_label)
            a.font = _font(9, True, AZUL)
            a.alignment = Alignment(vertical="top", wrap_text=True)
            b = ws.cell(row=row, column=3, value=t)
        else:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            b = ws.cell(row=row, column=2, value=t)
        b.font = _font(9)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        width = 100 if bold_label else 134
        ws.row_dimensions[row].height = max(15, 12.5 * math.ceil(len(t) / width) + 4)
        row += 1

    section("1. Para que serve esta planilha")
    text("O Normalizador IFC recebe os modelos exportados do Revit, Tekla e Civil 3D e os ajusta ao padrão "
         "Promon / BEP e à especificação IFC 4.3 da buildingSMART: classe de cada elemento, tipo, nomes e "
         "parâmetros PRM. A geometria nunca é alterada e os arquivos de origem nunca são modificados.")
    text("O que a ferramenta muda é definido por REGRAS. As regras desta planilha ainda são PROPOSTAS: foram "
         "montadas a partir dos modelos de exemplo e precisam ser validadas por quem conhece o BEP e a "
         "engenharia de cada disciplina. Nenhuma regra proposta é aplicada nos arquivos oficiais até ser aprovada.")

    section("2. Como ler uma regra")
    text("Toda regra tem a forma: SE o elemento atende a uma condição, ENTÃO a ferramenta muda algo nele.")
    text("Coluna “Aplica-se a”: quais elementos a regra alcança (programa de origem, disciplina, classe, nome…).",
         "Aplica-se a")
    text("Coluna “O que muda”: o que acontece com esses elementos (nova classe, tipo, parâmetros copiados…).",
         "O que muda")
    text("Uma peça real dos modelos de exemplo, com o valor antes e depois da regra. É o melhor ponto de "
         "partida para entender a regra.", "Exemplo real")
    text("Quantas peças a regra alcançou na simulação e em quantos arquivos.", "Peças afetadas / Arquivos")
    text("O que precisa ser respondido para aprovar a regra.", "Perguntas para validar")
    text("As últimas colunas (cinza) mostram a regra no formato técnico do arquivo de regras. Servem só "
         "para quem vai editar o arquivo; não é preciso analisá-las.", "Detalhe técnico")

    section("3. Como preencher")
    text("Leia o exemplo e responda às perguntas da regra.", "Passo 1")
    text("Escolha a decisão na coluna laranja “Decisão”: Aprovar · Ajustar · Rejeitar · Aguardar.", "Passo 2")
    text("Em “Ajustar” ou “Rejeitar”, explique em palavras simples na coluna “Comentário ou ajuste” "
         "(ex.: “diagonais da treliça da cobertura devem ser barra comum, não contraventamento”). "
         "Quem mantém a ferramenta traduz o comentário para o arquivo de regras.", "Passo 3")
    text("Indique o item do BEP que fundamenta a decisão, quando houver, e seu nome e a data.", "Passo 4")
    text("Na aba Geometria_Revisao, faça o mesmo para cada grupo de peças interpretado pela forma.", "Passo 5")

    section("4. Abas desta planilha")
    text("As regras propostas, agrupadas por tema, com explicação, exemplo e decisão.", "Regras_Validacao")
    text("Peças classificadas pela forma (vertical, horizontal, inclinada). É a parte mais sujeita a erro "
         "de interpretação: confira principalmente os grupos com muitas peças.", "Geometria_Revisao")
    text("Lista de valores da coluna Decisão (não editar).", "LISTAS")

    section("5. Glossário")
    for termo, definicao in GLOSSARIO:
        text(definicao, termo)

    section("6. Depois da validação")
    text("As decisões são levadas ao arquivo rules/normalizacao.json (ao lado do executável NormalizadorIFC.exe): "
         "regras aprovadas mudam para “aprovada” e recebem a referência do BEP; ajustes viram novas regras propostas, "
         "que podem ser simuladas de novo na própria ferramenta (opção “Incluir regras propostas”).")


def _rules_sheet(wb: Workbook, ruleset: RuleSet, usage: dict, n_reports: int, today: str):
    ws = wb.create_sheet("Regras_Validacao", 1)
    headers = ["Nº", "Regra", "Aplica-se a", "O que muda", "Exemplo real (antes → depois)", "Peças afetadas",
               "Arquivos", "Por que / alternativas", "Perguntas para validar",
               *DECISION_COLUMNS,
               "Situação atual", "Detalhe técnico – condição", "Detalhe técnico – resultado"]
    widths = [5, 13, 42, 40, 46, 10, 10, 40, 46, 12, 34, 18, 16, 11, 11, 36, 36]
    decision_from = headers.index(DECISION_COLUMNS[0])
    technical_from = headers.index("Situação atual")
    _sheet_top(
        ws, "REGRAS PROPOSTAS – VALIDAÇÃO COM O BEP",
        f"Regras versão {ruleset.version} · {len(ruleset.rules)} regras · simulação em {n_reports} arquivo(s) · {today}",
        len(headers),
        "Para cada regra: leia “Aplica-se a”, “O que muda” e o exemplo real; responda às perguntas e registre a "
        "decisão nas colunas laranja. As colunas cinza do fim são o formato técnico (não precisam ser analisadas). "
        "Instruções completas e glossário na aba LEIA-ME.",
    )
    _header(ws, 5, headers, widths, decision_from)
    dv = _decision_list(ws, len(DECISOES))

    themes: list[str] = []
    for r in ruleset.rules:
        t = r.get("tema") or "Outras"
        if t not in themes:
            themes.append(t)
    row, n = 6, 0
    for theme in themes:
        rules = [r for r in ruleset.rules if (r.get("tema") or "Outras") == theme]
        _band(ws, row, f"{theme}  ·  {len(rules)} regra(s)", len(headers), REFERENCIA, 10, color=AZUL, height=20)
        row += 1
        for i, r in enumerate(rules):
            n += 1
            u = usage.get(r["id"]) or {"casados": 0, "arquivos": 0, "exemplo": None, "nomes": Counter()}
            ex = u["exemplo"]
            exemplo = (f"Peça “{ex['nome']}” ({ex['arquivo']}):\n" + "\n".join(ex["linhas"])) if ex else \
                "Nenhuma peça alcançada na simulação."
            situacao = "proposta – em teste" if r["status"] == "proposta" else "aprovada"
            if uses_geometry(r["quando"]):
                situacao += " (geometria)"
            perguntas = "\n".join(f"• {p}" for p in r.get("perguntas", []))
            values = [n, r["id"], describe_applies_to(r["quando"]), describe_change(r["entao"]), exemplo,
                      u["casados"], u["arquivos"], r.get("justificativa", ""), perguntas,
                      "", "", r.get("avaliacao_bep", ""), "", "",
                      situacao, describe_condition(r["quando"]), describe_action(r["entao"])]
            _data_row(ws, row, values, widths, zebra=bool(i % 2), numeric={0, 5, 6},
                      decision_from=decision_from, technical_from=technical_from)
            dv.add(ws.cell(row=row, column=decision_from + 2))
            row += 1
    ws.freeze_panes = "E6"
    ws.auto_filter.ref = f"B5:{get_column_letter(len(headers) + 1)}{row - 1}"


def _geometry_sheet(wb: Workbook, groups: pd.DataFrame):
    ws = wb.create_sheet("Geometria_Revisao", 2)
    headers = ["Nº", "Nome da peça", "Classe na origem", "Forma detectada", "Qtd", "Comprimento típico (m)",
               "Ângulo com a vertical (°)", "O que a ferramenta faria", "Pergunta", *DECISION_COLUMNS,
               "Arquivos", "Exemplo – arquivo", "Exemplo – GlobalId (para localizar)"]
    widths = [5, 22, 22, 18, 7, 13, 12, 34, 50, 12, 34, 18, 16, 11, 10, 34, 26]
    decision_from = headers.index(DECISION_COLUMNS[0])
    technical_from = headers.index("Arquivos")
    _sheet_top(
        ws, "PEÇAS CLASSIFICADAS PELA FORMA – CONFERÊNCIA",
        f"{int(groups['qtd'].sum())} peças em {len(groups)} grupos · regras geométricas em teste (propostas)",
        len(headers),
        "Cada linha agrupa peças com o mesmo nome e a mesma forma. A ferramenta mediu a peça (vertical, horizontal, "
        "inclinada, plana) e propõe uma classe. Responda à pergunta e registre a decisão nas colunas laranja. "
        "O GlobalId de exemplo localiza a peça no visualizador.",
    )
    _header(ws, 5, headers, widths, decision_from)
    dv = _decision_list(ws, len(DECISOES))
    row = 6
    for i, g in enumerate(groups.itertuples()):
        acao, pergunta = geometry_texts(g)
        values = [i + 1, g.nome_normalizado or "(sem nome)", class_label(g.classe_original),
                  ORIENTACOES.get(g.orientacao, g.orientacao), int(g.qtd),
                  g.comprimento_mediano_m if pd.notna(g.comprimento_mediano_m) else None,
                  int(g.angulo_medio) if pd.notna(g.angulo_medio) else None,
                  acao, pergunta, "", "", "", "", "", int(g.arquivos), g.exemplo_arquivo, g.exemplo_global_id]
        _data_row(ws, row, values, widths, zebra=bool(i % 2), numeric={0, 4, 5, 6},
                  decision_from=decision_from, technical_from=technical_from)
        dv.add(ws.cell(row=row, column=decision_from + 2))
        row += 1
    ws.freeze_panes = "D6"
    ws.auto_filter.ref = f"B5:{get_column_letter(len(headers) + 1)}{row - 1}"


def build(ruleset: RuleSet, reports: list[Path], out_path: Path) -> Path:
    usage, geo = collect(reports)
    today = dt.date.today().strftime("%d/%m/%Y")
    wb = Workbook()
    listas = wb.active
    listas.title = "LISTAS"
    listas.sheet_view.showGridLines = False
    listas["A1"] = "DECISÃO"
    listas["A1"].font = _font(9, True, AZUL)
    for i, d in enumerate(DECISOES, start=2):
        listas.cell(row=i, column=1, value=d).font = _font()

    _readme(wb, ruleset, len(reports), today)
    _rules_sheet(wb, ruleset, usage, len(reports), today)
    groups = geometry_groups(geo)
    if not groups.empty:
        _geometry_sheet(wb, groups)
    wb.active = 0
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
